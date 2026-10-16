"""Pure film-stack generation and bounded, explicit table import.

All generated layers are in deposition order (substrate towards air).  The
builder only prepares a process; it neither saves a file nor evaluates optical
performance.  Thickness is physical thickness in nm, rate is nm/s, and a
quarter-wave candidate uses a user-supplied real refractive index at the stated
reference wavelength.  No material properties or machine positions are guessed.

Public contracts:
* ``build_film_stack(payload)`` returns ``{profile, summary}``.
* ``parse_film_stack_file(filename, file_bytes, **options)`` returns an internal
  table with columns, complete rows, source row numbers, and format metadata.
* ``film_stack_import_preview(table)`` exposes at most 12 table rows.
* ``preview_film_stack_import(filename, file_bytes, **options)`` composes these.
* ``film_stack_import_materials(table, column)`` returns bounded unique labels.
* ``film_stack_import_sheets(filename, file_bytes)`` lists XLSX sheet choices.
* ``build_imported_film_stack(table, payload)`` applies explicit column/unit/
  order mappings and returns the same builder result.

Common builder keys: mode, name, machine_file, substrate_file, sources (or legacy materials),
description?, notes?. Sources select a saved design by source_slot; alias and rate_nm_s
are optional and otherwise derive from its material name/filename and reference rate.
Legacy materials are explicit aliases only. Materials contain material, source_slot, rate_nm_s,
target_mode, monitor_slot?, mask_slots?, refractive_index?, and
refractive_index_source?. Explicit: layers[{material, thickness_nm}]. Periodic:
period_layers, repeat_count, prefix_layers?, suffix_layers?. Equal: sequence,
layer_count, total_thickness_nm. Quarter-wave: sequence, layer_count,
reference_wavelength_nm. Import: columns{material, thickness}, thickness_unit,
order, decimal_separator?. Column selectors are exact names or zero-based ints.
The process dict is intended for local storage; agent responses should return
the short summary rather than the expanded steps.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Mapping
from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
import math
from pathlib import Path
import re
import shlex
from typing import Any
from zipfile import BadZipFile, ZipFile

from .process_environment import validate_environment
from .process_profiles import (PROCESS_SCHEMA_VERSION, validate_process_profile,
                               source_selection_options)


MAX_FILM_STACK_LAYERS = 10_000
MAX_FILM_STACK_MATERIALS = 256
MAX_IMPORT_BYTES = 16 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_IMPORT_COLUMNS = 128
MAX_IMPORT_CELL_CHARS = 4096
MAX_IMPORT_HEADER_ROW = 1000
IMPORT_PREVIEW_ROWS = 12
SUPPORTED_FILM_STACK_SUFFIXES = (".csv", ".tsv", ".txt", ".xlsx", ".json")
THICKNESS_UNITS = {"nm": 1.0, "um": 1000.0, "angstrom": 0.1}
FILM_STACK_MODES = ("explicit", "periodic", "equal", "quarter_wave")
IMPORT_ORDERS = ("substrate_to_air", "air_to_substrate")
TF_CALC_EXPORT_REFERENCE = "https://tfcalc.com/files/mac_demo/manual.pdf"
MACLEOD_EXPORT_REFERENCE = "https://www.thinfilmcenter.com/support/changes.htm"


def _object(value: Any, label: str) -> Mapping:
    if not isinstance(value, Mapping):
        raise ValueError(f"{label}必须是对象")
    return value


def _no_extra(value: Mapping, allowed: set[str], label: str) -> None:
    extra = set(value) - allowed
    if extra:
        raise ValueError(f"{label}包含不支持的参数：{', '.join(sorted(map(str, extra)))}")


def _number(value: Any, label: str, *, positive: bool = True) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{label}必须是有限数值")
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(f"{label}必须是有限数值") from exc
    if not math.isfinite(result) or (positive and result <= 0):
        raise ValueError(f"{label}必须是大于 0 的有限数值")
    return result


def _integer(value: Any, label: str, maximum: int = MAX_FILM_STACK_LAYERS) -> int:
    if type(value) is not int or not 1 <= value <= maximum:
        raise ValueError(f"{label}必须为 1–{maximum} 的整数")
    return value


def _material(value: Any, label: str) -> str:
    if not isinstance(value, str) or not 1 <= len(value.strip()) <= 120:
        raise ValueError(f"{label}必须是 1–120 字的材料名")
    return value.strip()


def _materials(value: Any) -> dict[str, dict]:
    if not isinstance(value, list) or not 1 <= len(value) <= MAX_FILM_STACK_MATERIALS:
        raise ValueError(f"材料映射必须包含 1–{MAX_FILM_STACK_MATERIALS} 项")
    result = {}
    allowed = {"material", "source_slot", "rate_nm_s", "target_mode",
               "monitor_slot", "mask_slots", "refractive_index", "refractive_index_source", "environment"}
    for index, raw in enumerate(value, 1):
        entry = _object(raw, f"第 {index} 项材料映射")
        _no_extra(entry, allowed, f"第 {index} 项材料映射")
        name = _material(entry.get("material"), f"第 {index} 项材料名")
        if name in result:
            raise ValueError(f"源选择标签重复：{name}；同名材料或重复安装源请明确指定不同 alias")
        rate = _number(entry.get("rate_nm_s"), f"材料 {name} 沉积速率（nm/s）")
        source = _integer(entry.get("source_slot"), f"材料 {name} 蒸发源位次", 64)
        mode = entry.get("target_mode")
        if mode not in ("mean", "center", "monitor"):
            raise ValueError(f"材料 {name} 须明确选择 mean、center 或 monitor 厚度基准")
        normalized = {"material": name, "source_slot": source,
                      "rate_nm_s": rate, "target_mode": mode}
        if mode == "monitor":
            normalized["monitor_slot"] = _integer(entry.get("monitor_slot"),
                                                    f"材料 {name} 监测器位次", 64)
        elif entry.get("monitor_slot") is not None:
            raise ValueError(f"材料 {name} 仅 monitor 厚度基准允许指定 monitor_slot")
        masks = entry.get("mask_slots")
        if masks is not None:
            if (not isinstance(masks, list) or
                    any(type(slot) is not int or slot < 0 for slot in masks) or
                    len(set(masks)) != len(masks)):
                raise ValueError(f"材料 {name} 修正板须为不重复的非负整数位次列表；[] 表示不使用")
            normalized["mask_slots"] = sorted(masks)
        if entry.get("refractive_index") is not None:
            normalized["refractive_index"] = _number(entry["refractive_index"],
                                                       f"材料 {name} 折射率")
        if "refractive_index_source" in entry:
            source_note = entry["refractive_index_source"]
            if not isinstance(source_note, str) or len(source_note) > 1000:
                raise ValueError(f"材料 {name} 折射率来源须为不超过 1000 字的文字")
            normalized["refractive_index_source"] = source_note.strip()
        if entry.get("environment") is not None:
            normalized["environment"] = validate_environment(entry["environment"])
        result[name] = normalized
    return result


def _source_bindings(payload: Mapping, machine: Mapping | None) -> dict[str, dict]:
    """Resolve selections against trusted installed source snapshots only.

    Legacy ``materials`` entries are explicit aliases; with a machine context,
    they gain the same derived display metadata, but retain their supplied rate.
    This function performs no file I/O. Callers own authorized file resolution.
    """
    options = source_selection_options(machine)
    if "sources" in payload:
        if machine is None:
            raise ValueError("sources 选择需要已解析的机器与蒸发源文件；请通过工艺服务或编辑器生成")
        entries = []
        for index, selection in enumerate(payload["sources"], 1):
            slot = selection["source_slot"]
            if slot > len(options) or not options[slot - 1]["valid"]:
                raise ValueError(f"第 {index} 项选择的源位次 {slot} 没有有效的已保存蒸发源设计")
            source = options[slot - 1]
            alias = selection.get("alias") or source["material_name"] or source["file"]
            entries.append({**{key: value for key, value in selection.items() if key != "alias"},
                            "material": alias,
                            "rate_nm_s": selection.get("rate_nm_s", source["default_rate_nm_s"])})
        result = _materials(entries)
    else:
        result = _materials(payload.get("materials"))
    for name, entry in result.items():
        slot = entry["source_slot"]
        if machine is not None:
            if slot > len(options) or not options[slot - 1]["valid"]:
                raise ValueError(f"别名 {name} 选择的源位次 {slot} 没有有效的已保存蒸发源设计")
            option = options[slot - 1]
            entry.update(source_file=option["file"], material_name=option["material_name"],
                         material_id=option["material_id"])
            if entry.get("monitor_slot", 0) > len(machine.get("monitors", [])):
                raise ValueError(f"别名 {name} 监测器位次超出机器安装范围")
            if any(slot >= len(machine.get("masks", [])) for slot in entry.get("mask_slots", [])):
                raise ValueError(f"别名 {name} 修正板位次超出机器安装范围")
    return result


def _layers(value: Any, label: str, *, allow_empty: bool = False) -> list[dict]:
    minimum = 0 if allow_empty else 1
    if not isinstance(value, list) or not minimum <= len(value) <= MAX_FILM_STACK_LAYERS:
        raise ValueError(f"{label}须包含 {minimum}–{MAX_FILM_STACK_LAYERS} 层")
    result = []
    for index, item in enumerate(value, 1):
        raw = _object(item, f"{label}第 {index} 层")
        _no_extra(raw, {"material", "thickness_nm"}, f"{label}第 {index} 层")
        result.append({"material": _material(raw.get("material"), f"{label}第 {index} 层材料"),
                       "thickness_nm": _number(raw.get("thickness_nm"),
                                                f"{label}第 {index} 层物理厚度（nm）")})
    return result


def _sequence(value: Any) -> list[str]:
    if not isinstance(value, list) or not 1 <= len(value) <= MAX_FILM_STACK_MATERIALS:
        raise ValueError(f"材料顺序须包含 1–{MAX_FILM_STACK_MATERIALS} 个材料名")
    return [_material(item, "材料顺序中的材料名") for item in value]


_COMMON_KEYS = {"mode", "name", "machine_file", "substrate_file", "materials", "sources",
                "description", "notes", "environment"}
_MODE_KEYS = {
    "explicit": {"layers"},
    "periodic": {"period_layers", "repeat_count", "prefix_layers", "suffix_layers"},
    "equal": {"sequence", "layer_count", "total_thickness_nm"},
    "quarter_wave": {"sequence", "layer_count", "reference_wavelength_nm"},
    "import": {"columns", "thickness_unit", "order", "decimal_separator"},
}


def _assemble(payload: Mapping, layers: list[dict], materials: dict[str, dict],
              *, warnings: list[str] | None = None) -> dict:
    if not 1 <= len(layers) <= MAX_FILM_STACK_LAYERS:
        raise ValueError(f"膜系须包含 1–{MAX_FILM_STACK_LAYERS} 层")
    steps = []
    for index, layer in enumerate(layers, 1):
        name = layer["material"]
        if name not in materials:
            raise ValueError(f"第 {index} 层材料 {name} 尚未映射蒸发源与沉积速率")
        step = {key: deepcopy(value) for key, value in materials[name].items()
                if key in {"source_slot", "rate_nm_s", "target_mode", "monitor_slot", "mask_slots", "environment"}}
        step["target_nm"] = layer["thickness_nm"]
        steps.append(step)
    profile = validate_process_profile({
        "schema_version": PROCESS_SCHEMA_VERSION,
        **({"environment": payload["environment"]} if payload.get("environment") is not None else {}),
        "name": payload.get("name"), "machine_file": payload.get("machine_file"),
        "substrate_file": payload.get("substrate_file"), "steps": steps,
        "description": payload.get("description", ""), "notes": payload.get("notes", ""),
    })
    counts = Counter(layer["material"] for layer in layers)
    actual_materials = {materials[name].get("material_id") or materials[name].get("material_name")
                        or f"unassigned-source-{materials[name]['source_slot']}" for name in counts}
    summary = {"mode": payload["mode"], "layer_count": len(layers),
               "material_count": len(actual_materials) if "sources" in payload else len(counts),
               "source_count": len({materials[name]["source_slot"] for name in counts}),
               "total_thickness_nm": math.fsum(layer["thickness_nm"] for layer in layers),
               "nominal_duration_s": math.fsum(step["target_nm"] / step["rate_nm_s"]
                                               for step in profile["steps"]),
               "deposition_order": "substrate_to_air",
               "warnings": list(warnings or [])}
    summary["source_bindings"] = [{"alias": name, **{key: entry[key] for key in
        ("source_slot", "source_file", "material_name", "material_id", "rate_nm_s") if key in entry}}
        for name, entry in materials.items() if name in counts]
    if payload["mode"] == "quarter_wave":
        summary["reference_wavelength_nm"] = payload["reference_wavelength_nm"]
        summary["candidate_only"] = True
    return {"profile": profile, "summary": summary}


def build_film_stack(payload: Mapping, *, machine: Mapping | None = None) -> dict:
    """Expand an explicitly chosen rule into a validated v6 process dict."""
    payload = _object(payload, "膜系参数")
    from .film_stack_schema import film_stack_spec_data
    payload = film_stack_spec_data(payload)
    mode = payload.get("mode")
    if mode not in FILM_STACK_MODES:
        raise ValueError("膜系生成规则须为 explicit、periodic、equal 或 quarter_wave")
    _no_extra(payload, _COMMON_KEYS | _MODE_KEYS[mode], "膜系参数")
    materials = _source_bindings(payload, machine)
    warnings = []
    if mode == "explicit":
        layers = _layers(payload.get("layers"), "显式层表")
    elif mode == "periodic":
        period = _layers(payload.get("period_layers"), "重复周期")
        prefix = _layers(payload.get("prefix_layers", []), "前置层", allow_empty=True)
        suffix = _layers(payload.get("suffix_layers", []), "后置层", allow_empty=True)
        repeats = _integer(payload.get("repeat_count"), "重复次数")
        if len(prefix) + repeats * len(period) + len(suffix) > MAX_FILM_STACK_LAYERS:
            raise ValueError(f"展开后层数超过 {MAX_FILM_STACK_LAYERS} 层上限")
        layers = prefix + period * repeats + suffix
    else:
        sequence = _sequence(payload.get("sequence"))
        count = _integer(payload.get("layer_count"), "层数")
        if mode == "equal":
            thickness = _number(payload.get("total_thickness_nm"), "总物理厚度（nm）") / count
            layers = [{"material": sequence[index % len(sequence)], "thickness_nm": thickness}
                      for index in range(count)]
        else:
            wavelength = _number(payload.get("reference_wavelength_nm"), "参考波长（nm）")
            for name in sequence:
                if name not in materials or "refractive_index" not in materials[name]:
                    raise ValueError(f"四分之一波候选须明确提供材料 {name} 在参考波长的实数折射率")
            layers = [{"material": sequence[index % len(sequence)],
                       "thickness_nm": wavelength / (4 * materials[sequence[index % len(sequence)]]["refractive_index"])}
                      for index in range(count)]
            warnings.append("四分之一波规则仅生成法向入射、均匀透明膜层候选；未验证反射率、带宽、吸收或工艺可制造性。折射率须对应参考波长。")
    return _assemble(payload, layers, materials, warnings=warnings)


def _unique_json_object(pairs: list[tuple[str, Any]]) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"导入 JSON 参数重复：{key}")
        result[key] = value
    return result


def _cell(value: Any, label: str) -> Any:
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError(f"{label}包含非有限数值")
        return value
    if not isinstance(value, str):
        raise ValueError(f"{label}仅支持文字、数字或空值；请导出物理厚度值表")
    if len(value) > MAX_IMPORT_CELL_CHARS:
        raise ValueError(f"{label}超过 {MAX_IMPORT_CELL_CHARS} 字上限")
    return value.strip()


def _column_names(values: list[Any], has_header: bool) -> tuple[list[str], list[str]]:
    if not 2 <= len(values) <= MAX_IMPORT_COLUMNS:
        raise ValueError(f"导入表须包含 2–{MAX_IMPORT_COLUMNS} 列")
    columns, warnings, seen = [], [], set()
    for index, value in enumerate(values):
        name = str(value).strip() if has_header and value is not None else ""
        if not name:
            name = f"列 {index + 1}"
        if len(name) > 256:
            raise ValueError("导入表列名不得超过 256 字")
        if name in seen:
            original = name
            name = f"{name}（列 {index + 1}）"
            while name in seen:
                name += "*"
            warnings.append(f"重复列名 {original} 已显示为 {name}；请按列位次映射。")
        columns.append(name)
        seen.add(name)
    return columns, warnings


def _table(columns: list[str], rows: list[list[Any]], row_numbers: list[int], **metadata) -> dict:
    if not rows or len(rows) > MAX_FILM_STACK_LAYERS:
        raise ValueError(f"导入表须包含 1–{MAX_FILM_STACK_LAYERS} 条膜层记录")
    cleaned = []
    for row, number in zip(rows, row_numbers):
        if len(row) != len(columns):
            raise ValueError(f"导入表第 {number} 行列数与表头不一致；请重选分隔符或导出矩形表格")
        cleaned.append([_cell(value, f"第 {number} 行第 {index + 1} 列")
                        for index, value in enumerate(row)])
    hints = _column_hints(columns)
    return {"columns": columns, "rows": cleaned, "row_numbers": row_numbers,
            "row_count": len(cleaned), "column_suggestions": hints,
            "warnings": ["请确认所选列为物理厚度，明确厚度单位及层序，并为每种材料映射蒸发源、速率与厚度基准。", *metadata.pop("warnings", [])],
            **metadata}


def _column_hints(columns: list[str]) -> dict:
    """Suggest labels only; never interpret a file's units or deposition order."""
    material_index = thickness_index = unit = None
    for index, column in enumerate(columns):
        key = re.sub(r"[\s_()\[\]（）-]+", "", column.casefold())
        if material_index is None and key in {"material", "materialname", "材料", "膜料", "材料名"}:
            material_index = index
        if thickness_index is None and key in {"thicknessnm", "targetnm", "physicalthickness", "physicalthicknessnm", "物理厚度", "物理厚度nm", "厚度nm", "thickness"}:
            thickness_index = index
            if key.endswith("nm"):
                unit = "nm"
    return {"material": material_index, "thickness": thickness_index, "thickness_unit": unit}


def _text_table(filename: str, content: bytes, *, header_row: int, delimiter: str | None,
                encoding: str, has_header: bool) -> dict:
    allowed_encodings = {"utf-8", "utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "gb18030"}
    if encoding not in allowed_encodings:
        raise ValueError("文本编码仅支持 utf-8(-sig)、utf-16(-le/-be) 或 gb18030；请明确选择编码")
    try:
        text = content.decode(encoding)
    except UnicodeError as exc:
        raise ValueError("文件不是所选编码的文本；请选择正确编码，或从原软件重新导出 UTF-8 CSV/TSV") from exc
    if "\x00" in text:
        raise ValueError("检测到二进制内容；不支持专有二进制膜系文件，请导出 CSV/TSV/物理厚度文本表")
    suffix = Path(filename).suffix.casefold()
    if delimiter is None:
        if suffix == ".csv":
            delimiter = ","
        elif suffix == ".tsv":
            delimiter = "\t"
        else:
            sample_lines = [line for line in text.splitlines()[header_row - 1:header_row + 20] if line.strip()]
            candidates = [separator for separator in ("\t", ",", ";")
                          if sample_lines and all(separator in line for line in sample_lines)]
            if len(candidates) != 1:
                raise ValueError("TXT 分隔符不明确；请明确选择逗号、制表符、分号或 whitespace（空白）")
            delimiter = candidates[0]
    if delimiter not in {",", "\t", ";", "whitespace"}:
        raise ValueError("分隔符须为逗号、制表符、分号或 whitespace")
    parsed, numbers = [], []
    try:
        if delimiter == "whitespace":
            entries = ((number, shlex.split(line, comments=False, posix=True))
                       for number, line in enumerate(text.splitlines(), 1) if line.strip())
        else:
            reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
            entries = ((reader.line_num, row) for row in reader)
        for number, row in entries:
            if number < header_row or not row or all(not str(value).strip() for value in row):
                continue
            if len(parsed) > MAX_FILM_STACK_LAYERS:
                raise ValueError(f"导入表超过 {MAX_FILM_STACK_LAYERS} 条膜层记录")
            parsed.append(row)
            numbers.append(number)
    except (csv.Error, ValueError) as exc:
        raise ValueError(f"无法解析文本表：{exc}") from exc
    if not parsed:
        raise ValueError("所选表头行之后没有表格内容")
    columns, warnings = _column_names(parsed[0], has_header)
    if has_header:
        parsed, numbers = parsed[1:], numbers[1:]
    return _table(columns, parsed, numbers, filename=filename, format=suffix.lstrip("."),
                  sheet=None, sheets=[], delimiter=delimiter, header_row=header_row,
                  has_header=has_header, encoding=encoding, warnings=warnings)


def _json_table(filename: str, content: bytes) -> dict:
    try:
        raw = json.loads(content.decode("utf-8-sig"), object_pairs_hook=_unique_json_object,
                         parse_constant=lambda token: (_ for _ in ()).throw(ValueError(f"不支持 JSON 非有限数值 {token}")))
    except (UnicodeError, json.JSONDecodeError, RecursionError) as exc:
        raise ValueError(f"无法解析 UTF-8 JSON 膜层表：{exc}") from exc
    declarations = {}
    if isinstance(raw, Mapping):
        for key in ("thickness_unit", "order"):
            if key in raw:
                declarations[key] = raw[key]
        if "columns" in raw and "rows" in raw:
            _no_extra(raw, {"columns", "rows", "thickness_unit", "order"}, "JSON 表格")
            if not isinstance(raw["columns"], list) or not isinstance(raw["rows"], list):
                raise ValueError("JSON columns 与 rows 须为列表")
            if any(not isinstance(column, str) or not column.strip() for column in raw["columns"]):
                raise ValueError("JSON columns 须为非空文字列名")
            columns, warnings = _column_names(raw["columns"], True)
            if any(not isinstance(row, list) for row in raw["rows"]):
                raise ValueError("JSON rows 中每一行须为列表")
            return _table(columns, raw["rows"], list(range(1, len(raw["rows"]) + 1)),
                          filename=filename, format="json", sheet=None, sheets=[],
                          declared_metadata=declarations, warnings=warnings)
        _no_extra(raw, {"layers", "thickness_unit", "order"}, "JSON 膜层表")
        raw = raw.get("layers")
    if not isinstance(raw, list) or not raw or len(raw) > MAX_FILM_STACK_LAYERS:
        raise ValueError(f"JSON 须为 1–{MAX_FILM_STACK_LAYERS} 个膜层对象，或包含 layers 的对象")
    if any(not isinstance(row, Mapping) for row in raw):
        raise ValueError("JSON 每层须为包含材料与物理厚度的对象")
    original_columns = list(dict.fromkeys(key for row in raw for key in row))
    columns, warnings = _column_names(original_columns, True)
    rows = [[row.get(column) for column in original_columns] for row in raw]
    return _table(columns, rows, list(range(1, len(rows) + 1)), filename=filename,
                  format="json", sheet=None, sheets=[], declared_metadata=declarations, warnings=warnings)


def _xlsx_workbook(content: bytes):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("XLSX 解压后过大；请导出仅包含膜层表的工作簿或 CSV")
            if "xl/workbook.xml" not in archive.namelist() or any(entry.flag_bits & 1 for entry in entries):
                raise ValueError("仅支持未加密的标准 XLSX 工作簿")
    except (BadZipFile, OSError) as exc:
        raise ValueError("文件不是标准 XLSX；旧版 XLS、宏或专有二进制文件请另存为 XLSX/CSV") from exc
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX 导入组件 openpyxl 未安装；请安装项目依赖或另存为 CSV/TSV") from exc
    try:
        return load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError(f"无法读取标准 XLSX 工作簿：{exc}") from exc


def film_stack_import_sheets(filename: str, file_bytes: bytes) -> list[str]:
    """List workbook choices without selecting or combining their layer tables."""
    if not isinstance(filename, str) or Path(filename).suffix.casefold() != ".xlsx":
        return []
    if not isinstance(file_bytes, bytes) or not 0 < len(file_bytes) <= MAX_IMPORT_BYTES:
        raise ValueError(f"导入文件须非空且不超过 {MAX_IMPORT_BYTES // (1024 * 1024)} MiB")
    workbook = _xlsx_workbook(file_bytes)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _xlsx_table(filename: str, content: bytes, *, sheet: str | None,
                header_row: int, has_header: bool) -> dict:
    workbook = _xlsx_workbook(content)
    try:
        names = workbook.sheetnames
        if sheet is None:
            if len(names) != 1:
                raise ValueError("工作簿含多个工作表；请明确选择 sheet，不能自动合并")
            selected = names[0]
        elif not isinstance(sheet, str) or sheet not in names:
            raise ValueError("所选 XLSX 工作表不存在")
        else:
            selected = sheet
        worksheet = workbook[selected]
        if (worksheet.max_column or 0) > MAX_IMPORT_COLUMNS:
            raise ValueError(f"XLSX 超过 {MAX_IMPORT_COLUMNS} 列；请只导出膜层表")
        reported_rows = worksheet.max_row
        if (reported_rows or 0) > header_row + MAX_FILM_STACK_LAYERS:
            raise ValueError(f"XLSX 超过 {MAX_FILM_STACK_LAYERS} 条膜层记录；请移除多余空行和格式")
        # Incorrect producer dimensions must not silently truncate imported data.
        worksheet.reset_dimensions()
        parsed, numbers = [], []
        max_scan_row = header_row + MAX_FILM_STACK_LAYERS + 1
        for number, row in enumerate(worksheet.iter_rows(min_row=header_row,
                                max_row=max_scan_row, max_col=MAX_IMPORT_COLUMNS + 1), header_row):
            values = [cell.value for cell in row]
            if all(value is None for value in values):
                continue
            while values and values[-1] is None:
                values.pop()
            if len(values) > MAX_IMPORT_COLUMNS:
                raise ValueError(f"XLSX 超过 {MAX_IMPORT_COLUMNS} 列")
            if any(cell.data_type == "f" for cell in row):
                raise ValueError(f"XLSX 第 {number} 行包含公式；请在原表复制并粘贴为值后导入，避免缓存值失真")
            parsed.append(values)
            numbers.append(number)
        if not parsed:
            raise ValueError("所选 XLSX 工作表没有膜层表数据")
        width = max(len(row) for row in parsed)
        parsed = [row + [None] * (width - len(row)) for row in parsed]
        columns, warnings = _column_names(parsed[0], has_header)
        if has_header:
            parsed, numbers = parsed[1:], numbers[1:]
        return _table(columns, parsed, numbers, filename=filename, format="xlsx", sheet=selected,
                      sheets=names, header_row=header_row, has_header=has_header, warnings=warnings)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"无法读取标准 XLSX 膜层表：{exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def parse_film_stack_file(filename: str, file_bytes: bytes, *, sheet: str | None = None,
                          header_row: int = 1, delimiter: str | None = None,
                          encoding: str = "utf-8-sig", has_header: bool = True) -> dict:
    """Read a complete bounded table; no physical interpretation occurs here.

    TFCalc's material/physical-thickness text export and Macleod's CSV/TSV
    tables use the generic table reader. Proprietary native formats, spectra,
    QWOT-only tables, and binary files are not decoded by this module.
    """
    if not isinstance(filename, str) or not filename.strip() or len(filename) > 4096:
        raise ValueError("导入文件名须为非空文字")
    suffix = Path(filename).suffix.casefold()
    if suffix not in SUPPORTED_FILM_STACK_SUFFIXES:
        raise ValueError("不支持该专有或二进制膜系格式；请从 TFCalc/Macleod 或原软件导出物理厚度 CSV/TSV/TXT，或另存为 XLSX/JSON 表")
    if not isinstance(file_bytes, bytes) or not 0 < len(file_bytes) <= MAX_IMPORT_BYTES:
        raise ValueError(f"导入文件须非空且不超过 {MAX_IMPORT_BYTES // (1024 * 1024)} MiB")
    _integer(header_row, "表头/数据起始行", MAX_IMPORT_HEADER_ROW)
    if type(has_header) is not bool:
        raise ValueError("has_header 必须为布尔值；无表头的 TFCalc 两列表请选择 false")
    if suffix == ".json":
        if sheet is not None or header_row != 1 or delimiter is not None or not has_header:
            raise ValueError("JSON 对象层表不使用工作表、起始行、分隔符或无表头选项")
        return _json_table(filename, file_bytes)
    if suffix == ".xlsx":
        if delimiter is not None:
            raise ValueError("XLSX 不使用文本分隔符")
        return _xlsx_table(filename, file_bytes, sheet=sheet, header_row=header_row, has_header=has_header)
    if sheet is not None:
        raise ValueError("仅 XLSX 使用 sheet 工作表选项")
    return _text_table(filename, file_bytes, header_row=header_row, delimiter=delimiter,
                       encoding=encoding, has_header=has_header)


def film_stack_import_preview(table: Mapping) -> dict:
    """Small UI/agent preview. Complete imported rows stay local to the service."""
    table = _object(table, "导入表")
    return {key: deepcopy(table[key]) for key in (
        "columns", "row_count", "column_suggestions", "warnings", "filename", "format",
        "sheet", "sheets", "declared_metadata", "delimiter", "header_row", "has_header", "encoding")
        if key in table} | {"preview_rows": deepcopy(table["rows"][:IMPORT_PREVIEW_ROWS]),
                            "preview_row_numbers": list(table["row_numbers"][:IMPORT_PREVIEW_ROWS]),
                            "preview_truncated": table["row_count"] > IMPORT_PREVIEW_ROWS}


def preview_film_stack_import(filename: str, file_bytes: bytes, **options) -> dict:
    return film_stack_import_preview(parse_film_stack_file(filename, file_bytes, **options))


def _column_index(columns: list[str], selector: Any, label: str) -> int:
    if type(selector) is int and 0 <= selector < len(columns):
        return selector
    if isinstance(selector, str) and selector in columns:
        return columns.index(selector)
    raise ValueError(f"{label}须为已存在的列名或从 0 开始的列位次")


def film_stack_import_materials(table: Mapping, column: str | int) -> list[str]:
    """Unique exact labels in table order; case-sensitive, without substitutions."""
    index = _column_index(table["columns"], column, "材料列")
    result, seen = [], set()
    for row, number in zip(table["rows"], table["row_numbers"]):
        name = _material(row[index], f"导入第 {number} 行材料")
        if name not in seen:
            result.append(name)
            seen.add(name)
        if len(result) > MAX_FILM_STACK_MATERIALS:
            raise ValueError(f"导入表超过 {MAX_FILM_STACK_MATERIALS} 种材料")
    return result


def build_imported_film_stack(table: Mapping, payload: Mapping, *, machine: Mapping | None = None) -> dict:
    """Map physical layer data into deposition steps using explicit user choices."""
    table, payload = _object(table, "导入表"), _object(payload, "导入映射")
    from .film_stack_schema import film_stack_spec_data
    payload = film_stack_spec_data(payload)
    if payload.get("mode") != "import":
        raise ValueError("导入映射 mode 须为 import")
    _no_extra(payload, _COMMON_KEYS | _MODE_KEYS["import"], "导入映射")
    columns = _object(payload.get("columns"), "导入列映射")
    _no_extra(columns, {"material", "thickness"}, "导入列映射")
    material_column = _column_index(table["columns"], columns.get("material"), "材料列")
    thickness_column = _column_index(table["columns"], columns.get("thickness"), "物理厚度列")
    if material_column == thickness_column:
        raise ValueError("材料列和物理厚度列不能是同一列")
    thickness_name = str(table["columns"][thickness_column]).casefold()
    if any(term in thickness_name for term in ("qwot", "optical", "光学厚度", "quarter wave", "quarter-wave")):
        raise ValueError("所选列属于光学厚度/QWOT；请先在原软件导出物理厚度，不能直接当作 nm")
    unit = payload.get("thickness_unit")
    if unit not in THICKNESS_UNITS:
        raise ValueError("须明确指定物理厚度单位 thickness_unit=nm、um 或 angstrom")
    order = payload.get("order")
    if order not in IMPORT_ORDERS:
        raise ValueError("须明确指定层序 order=substrate_to_air 或 air_to_substrate")
    decimal = payload.get("decimal_separator", ".")
    if decimal not in (".", ","):
        raise ValueError("小数符号 decimal_separator 仅支持 . 或 ,；不支持分组千位符")
    declarations = table.get("declared_metadata", {})
    if declarations.get("thickness_unit") is not None and declarations["thickness_unit"] != unit:
        raise ValueError("所选厚度单位与 JSON 声明不一致；请修正文件或映射")
    if declarations.get("order") is not None and declarations["order"] != order:
        raise ValueError("所选层序与 JSON 声明不一致；请修正文件或映射")
    rows, row_numbers = table.get("rows"), table.get("row_numbers")
    if (not isinstance(rows, list) or not 1 <= len(rows) <= MAX_FILM_STACK_LAYERS or
            not isinstance(row_numbers, list) or len(rows) != len(row_numbers)):
        raise ValueError("导入表记录数或来源行号无效；请重新解析文件")
    materials, layers = _source_bindings(payload, machine), []
    for row, number in zip(rows, row_numbers):
        if not isinstance(row, list) or len(row) != len(table["columns"]):
            raise ValueError(f"导入第 {number} 行列数无效")
        name = _material(row[material_column], f"导入第 {number} 行材料")
        thickness = row[thickness_column]
        if isinstance(thickness, str):
            if decimal == ",":
                if "." in thickness:
                    raise ValueError(f"导入第 {number} 行厚度与所选逗号小数符号不一致")
                thickness = thickness.replace(",", ".")
            elif "," in thickness:
                raise ValueError(f"导入第 {number} 行厚度含逗号；请明确选择小数符号或移除千位符")
        value = _number(thickness, f"导入第 {number} 行物理厚度") * THICKNESS_UNITS[unit]
        layers.append({"material": name, "thickness_nm": _number(value, f"导入第 {number} 行换算厚度（nm）")})
    if order == "air_to_substrate":
        layers.reverse()
    warnings = list(table.get("warnings", []))
    warnings.append("仅导入物理膜层厚度；外部软件的色散、光谱指标、监测算法及优化约束没有转换为工艺参数。")
    result = _assemble(payload, layers, materials, warnings=warnings)
    result["summary"]["import_format"] = table.get("format")
    result["summary"]["input_order"] = order
    result["summary"]["input_thickness_unit"] = unit
    return result
